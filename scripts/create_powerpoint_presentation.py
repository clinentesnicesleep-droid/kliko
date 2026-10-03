# -*- coding: utf-8 -*-
"""
Genera la presentación oficial en PowerPoint (.pptx) en formato panorámico 16:9
del III Plan Municipal de Juventud de Orcera (2027-2031) para el Alcalde y el Pleno,
incorporando las 7 diapositivas individuales de cronograma quinquenal (una por cada Eje)
con todas sus acciones y duración en años (1, 2, 3, 4, 5 / 2027-2031).
"""

import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Colores corporativos del III Plan
    DARK_BG = RGBColor(15, 23, 42)       # Slate 900
    PINE_GREEN = RGBColor(6, 78, 59)     # Verde Pino Orcera
    EMERALD = RGBColor(16, 185, 129)     # Verde Esmeralda
    AMURJO_CYAN = RGBColor(6, 182, 212)  # Azul Amurjo
    CORAL = RGBColor(255, 87, 34)        # Naranja Coral
    WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800

    EJE_COLORS = {
        1: (RGBColor(16, 185, 129), RGBColor(6, 78, 59)),    # Emerald
        2: (RGBColor(245, 158, 11), RGBColor(120, 53, 15)),  # Amber
        3: (RGBColor(6, 182, 212), RGBColor(12, 74, 110)),   # Cyan
        4: (RGBColor(139, 92, 246), RGBColor(76, 29, 149)),  # Purple
        5: (RGBColor(236, 72, 153), RGBColor(131, 24, 67)),  # Pink
        6: (RGBColor(59, 130, 246), RGBColor(30, 58, 138)),  # Blue
        7: (RGBColor(20, 184, 166), RGBColor(17, 94, 89))    # Teal
    }

    img_dir = os.path.join(os.path.dirname(__file__), '..', 'img')
    hero_img = os.path.join(img_dir, 'hero-orcera-youth.jpg')
    amurjo_img = os.path.join(img_dir, 'amurjo-pool.jpg')
    asoc_img = os.path.join(img_dir, 'asociacion-juvenil.jpg')

    # Pantallazos móviles reales
    m_inicio = os.path.join(img_dir, 'pantallazo_movil_inicio.png')
    m_ejes = os.path.join(img_dir, 'pantallazo_movil_ejes.png')
    m_gamif = os.path.join(img_dir, 'pantallazo_movil_gamificacion.png')
    m_buzon = os.path.join(img_dir, 'pantallazo_movil_buzon.png')
    m_promo = os.path.join(img_dir, 'pantallazo_movil_promocion.png')

    # Helper function to add background color
    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    # Helper function to add a header on content slides
    def add_slide_header(slide, tag_text, title_text, tag_color=AMURJO_CYAN):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tag = tf.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_color
        p_tag.space_after = Pt(2)

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # =========================================================================
    # DIAPOSITIVA 1: PORTADA INSTITUCIONAL
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_BG)

    if os.path.exists(hero_img):
        slide1.shapes.add_picture(hero_img, Inches(6.8), Inches(0), Inches(6.533), Inches(7.5))
        overlay = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(0), Inches(1.5), Inches(7.5))
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = DARK_BG
        overlay.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(6.0), Inches(5.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "🏛️ AYUNTAMIENTO DE ORCERA · SIERRA DE SEGURA"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "III Plan Municipal\\nde Juventud"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Periodo de Vigencia Quinquenal 2027–2031"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = AMURJO_CYAN
    p.space_after = Pt(18)

    p = tf1.add_paragraph()
    p.text = "Presentación Oficial ante el Alcalde y el Pleno Municipal:\nPropuesta de Aprobación, Aplicación Web Cívica, Cronograma Quinquenal de los 7 Ejes y Estrategia de Dinamización Juvenil."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(24)

    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(5.2), Inches(0.9))
    badge.fill.solid()
    badge.fill.fore_color.rgb = PINE_GREEN
    badge.line.color.rgb = EMERALD
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "✓ 7 Ejes Estratégicos · 63 Acciones Oficiales · 75.000 € Inversión"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p = tf_b.add_paragraph()
    p.text = "✓ App PWA Móvil Activa · Premios Amurjo · Asociación Juvenil"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(209, 250, 229)

    # =========================================================================
    # DIAPOSITIVA 2: EL DESAFÍO Y DIAGNÓSTICO
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, DARK_BG)
    add_slide_header(slide2, "1. Diagnóstico de Partida", "El Desafío: Conectar con la Juventud y Frenar la Despoblación Rural")

    # Tarjeta Izquierda (Puntos Críticos)
    box_izq = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    box_izq.fill.solid()
    box_izq.fill.fore_color.rgb = CARD_BG
    box_izq.line.color.rgb = RGBColor(239, 68, 68)

    tf_izq = box_izq.text_frame
    tf_izq.word_wrap = True
    p = tf_izq.paragraphs[0]
    p.text = "SITUACIÓN ANTERIOR / PROBLEMAS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(248, 113, 113)
    p.space_after = Pt(12)

    puntos_problemas = [
        "Falta de canales directos de participación juvenil en las decisiones municipales.",
        "Riesgo de despoblación juvenil y desconexión con los recursos locales de Orcera.",
        "Desconocimiento generalizado de las ayudas, cursos y actividades municipales.",
        "Falta de transparencia en presupuestos y plazos de ejecución.",
        "Asociación juvenil histórica inactiva o con baja implicación de nuevas generaciones."
    ]
    for pt in puntos_problemas:
        p = tf_izq.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = WHITE
        p.space_after = Pt(8)

    # Tarjeta Derecha (Soluciones)
    box_der = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1))
    box_der.fill.solid()
    box_der.fill.fore_color.rgb = CARD_BG
    box_der.line.color.rgb = EMERALD

    tf_der = box_der.text_frame
    tf_der.word_wrap = True
    p = tf_der.paragraphs[0]
    p.text = "SOLUCIÓN: EL III PLAN Y SU APP DIGITAL"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(12)

    puntos_soluciones = [
        "Plataforma Digital Cívica accesible 24/7 desde cualquier teléfono móvil sin descargas.",
        "7 Ejes estratégicos con 63 acciones detalladas con responsables y calendario.",
        "Sistema de Gamificación 'Premios Amurjo': pases de piscina y pádel por participar.",
        "Reactivación de la Asociación Juvenil con Carnet Digital Oficial y Cuota Cero.",
        "Buzón Participativo Directo con traslado de propuestas al Pleno Municipal."
    ]
    for pt in puntos_soluciones:
        p = tf_der.add_paragraph()
        p.text = f"✓ {pt}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = WHITE
        p.space_after = Pt(8)

    # =========================================================================
    # DIAPOSITIVA 3: CÓMO VEN LA APP LOS JÓVENES EN SUS TELÉFONOS
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, DARK_BG)
    add_slide_header(slide3, "2. Experiencia de Usuario", "Cómo Ven la App los Jóvenes en sus Propios Teléfonos Móviles")

    # Tarjeta explicativa
    box_info3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(8.3), Inches(5.1))
    box_info3.fill.solid()
    box_info3.fill.fore_color.rgb = CARD_BG
    box_info3.line.color.rgb = AMURJO_CYAN

    tf3 = box_info3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "DISEÑADA ESPECÍFICAMENTE PARA LA JUVENTUD RURAL ACTUAL"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = AMURJO_CYAN
    p.space_after = Pt(12)

    features_movil = [
        ("Interfaz Móvil Nativa (PWA):", "No requiere buscar en Google Play ni App Store; se instala en 1 segundo con icono directo en pantalla."),
        ("Navegación Visual Intuitiva:", "Colores vibrantes, tipografía moderna e iconos por cada uno de los 7 ejes y sus proyectos."),
        ("Consulta Rápida de Proyectos:", "Los jóvenes ven exactamente qué actividades hay, en qué estado están y qué concejalía responde."),
        ("Votación y Participación Ágil:", "Encuestas, propuestas y valoración de acciones en menos de 3 toques."),
        ("Puntos Cívicos Acumulables:", "Recompensas automáticas canjeables por pases y ventajas en instalaciones locales.")
    ]
    for title_f, desc_f in features_movil:
        p = tf3.add_paragraph()
        p.text = f"• {title_f} "
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = WHITE
        r = p.add_run()
        r.text = desc_f
        r.font.bold = False
        r.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # Imagen móvil a la derecha
    if os.path.exists(m_inicio):
        slide3.shapes.add_picture(m_inicio, Inches(9.5), Inches(1.8), Inches(2.9), Inches(5.1))

    # =========================================================================
    # DIAPOSITIVAS 4 A 10: LAS 7 TRANSPARENCIAS DE CRONOGRAMA POR CADA EJE (9 ACCIONES C/U)
    # =========================================================================
    with open(os.path.join(os.path.dirname(__file__), 'ejes_extracted.json'), encoding='utf-8') as f:
        ejes = json.load(f)

    from check_curated_titles import curated_titles, eje_meta

    for eje in ejes:
        num = eje['numero']
        meta = eje_meta[num]
        color_rgb, dark_accent = EJE_COLORS[num]
        icono = meta['icono']
        presupuesto = meta['presupuesto']
        subtitulo = meta['subtitulo']

        slide_eje = prs.slides.add_slide(blank_layout)
        set_slide_background(slide_eje, DARK_BG)
        add_slide_header(
            slide_eje,
            f"3.{num}. Planificación Temporal · Eje {num} de 7",
            f"{icono} Eje {num}: {eje['titulo']} (9 Acciones · 2027–2031)",
            tag_color=color_rgb
        )

        # Barra subheader con presupuesto y resumen de acciones
        sub_box = slide_eje.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.42))
        sub_box.fill.solid()
        sub_box.fill.fore_color.rgb = CARD_BG
        sub_box.line.color.rgb = color_rgb
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = f"  Presupuesto Quinquenal: {presupuesto}   |   {subtitulo}   |   Leyenda: [■] Año Activo  [—] No Activo"
        p_sub.font.size = Pt(8.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = WHITE

        # Tabla Gantt: 10 filas (1 cabecera + 9 acciones), 7 columnas
        rows_cnt = 10
        cols_cnt = 7
        t_left, t_top, t_width, t_height = Inches(0.8), Inches(2.0), Inches(11.733), Inches(5.1)
        tbl_shape = slide_eje.shapes.add_table(rows_cnt, cols_cnt, t_left, t_top, t_width, t_height)
        tbl = tbl_shape.table

        # Anchos de columna
        tbl.columns[0].width = Inches(5.833) # Código + Título
        tbl.columns[1].width = Inches(0.92)  # 2027
        tbl.columns[2].width = Inches(0.92)  # 2028
        tbl.columns[3].width = Inches(0.92)  # 2029
        tbl.columns[4].width = Inches(0.92)  # 2030
        tbl.columns[5].width = Inches(0.92)  # 2031
        tbl.columns[6].width = Inches(1.30)  # Duración

        # Fila Cabecera
        tbl_headers = ["Medida / Proyecto Oficial (9)", "2027 (A1)", "2028 (A2)", "2029 (A3)", "2030 (A4)", "2031 (A5)", "Duración Total"]
        for c_idx, h_text in enumerate(tbl_headers):
            c = tbl.cell(0, c_idx)
            c.fill.solid()
            c.fill.fore_color.rgb = dark_accent
            p = c.text_frame.paragraphs[0]
            p.text = h_text
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER

        # Filas de Acciones (1 a 9)
        for r_idx, a in enumerate(eje['acciones'], start=1):
            code = a['codigo']
            title = curated_titles.get(code, a['titulo'])
            years = a['anos']
            dur_len = len(years)

            # Celda Medida / Proyecto
            c_med = tbl.cell(r_idx, 0)
            c_med.fill.solid()
            c_med.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else RGBColor(22, 30, 46)
            p_m = c_med.text_frame.paragraphs[0]
            p_m.text = f"[{code}] {title}"
            p_m.font.size = Pt(8.2)
            p_m.font.bold = True
            p_m.font.color.rgb = WHITE

            # Celdas de Años (1 a 5)
            for y_offset, yr in enumerate([2027, 2028, 2029, 2030, 2031]):
                c_yr = tbl.cell(r_idx, y_offset + 1)
                c_yr.fill.solid()
                if yr in years:
                    c_yr.fill.fore_color.rgb = dark_accent
                    p_y = c_yr.text_frame.paragraphs[0]
                    p_y.text = "████"
                    p_y.font.size = Pt(8)
                    p_y.font.bold = True
                    p_y.font.color.rgb = color_rgb
                    p_y.alignment = PP_ALIGN.CENTER
                else:
                    c_yr.fill.fore_color.rgb = RGBColor(18, 25, 38)
                    p_y = c_yr.text_frame.paragraphs[0]
                    p_y.text = "—"
                    p_y.font.size = Pt(8)
                    p_y.font.color.rgb = TEXT_MUTED
                    p_y.alignment = PP_ALIGN.CENTER

            # Celda Duración
            c_dur = tbl.cell(r_idx, 6)
            c_dur.fill.solid()
            c_dur.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else RGBColor(22, 30, 46)
            p_d = c_dur.text_frame.paragraphs[0]
            
            if dur_len == 5:
                p_d.text = "5 Años (100%)"
                p_d.font.color.rgb = EMERALD
            elif dur_len == 4:
                p_d.text = "4 Años (80%)"
                p_d.font.color.rgb = AMURJO_CYAN
            elif dur_len == 3:
                p_d.text = "3 Años (60%)"
                p_d.font.color.rgb = RGBColor(96, 165, 250)
            elif dur_len == 2:
                p_d.text = "2 Años (40%)"
                p_d.font.color.rgb = RGBColor(192, 132, 252)
            else:
                p_d.text = f"1 Año ({years[0]})"
                p_d.font.color.rgb = RGBColor(251, 191, 36)

            p_d.font.size = Pt(8)
            p_d.font.bold = True
            p_d.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # DIAPOSITIVA 11: LOS 7 EJES Y RIGOR PRESUPUESTARIO
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, DARK_BG)
    add_slide_header(slide4, "4. Rigor Presupuestario", "Inversión Responsable: 75.000 € (15.000 € / Año en 7 Ejes)")

    # Caja Izquierda: Desglose
    box_ejes = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(8.3), Inches(5.3))
    box_ejes.fill.solid()
    box_ejes.fill.fore_color.rgb = CARD_BG
    box_ejes.line.color.rgb = RGBColor(51, 65, 85)

    tf_e = box_ejes.text_frame
    tf_e.word_wrap = True
    p = tf_e.paragraphs[0]
    p.text = "DESGLOSE QUINQUENAL Y GARANTÍAS PARA EL AYUNTAMIENTO"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = AMURJO_CYAN
    p.space_after = Pt(12)

    ejes_data = [
        ("Eje 1: Entorno Sostenible, Digital y Conectado", "12.500 €", "16.7%"),
        ("Eje 2: Ocio Saludable, Cultura, Deporte y Creación", "13.000 €", "17.3%"),
        ("Eje 3: Emancipación, Empleo, Formación y Vivienda", "14.500 €", "19.3%"),
        ("Eje 4: Participación Juvenil y Voluntariado", "10.000 €", "13.3%"),
        ("Eje 5: Salud, Bienestar Emocional e Inclusión", "9.500 €", "12.7%"),
        ("Eje 6: Igualdad, Diversidad y Convivencia", "8.000 €", "10.7%"),
        ("Eje 7: Gobernanza, Innovación y Evaluación", "7.500 €", "10.0%")
    ]
    for eje, importe, pct in ejes_data:
        p = tf_e.add_paragraph()
        p.text = f"• {eje}: "
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        r1 = p.add_run()
        r1.text = f"{importe} "
        r1.font.bold = True
        r1.font.color.rgb = EMERALD
        r2 = p.add_run()
        r2.text = f"({pct})"
        r2.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(3)

    p = tf_e.add_paragraph()
    p.text = "\n✓ Control por facturas asociadas a proveedores oficiales"
    p.font.size = Pt(9.5)
    p.font.color.rgb = WHITE
    p = tf_e.add_paragraph()
    p.text = "✓ Valoración ciudadana en tiempo real (satisfacción actual: 4.8 ★)"
    p.font.size = Pt(9.5)
    p.font.color.rgb = WHITE
    p = tf_e.add_paragraph()
    p.text = "✓ Datos exportables para justificación ante el IAJ, Diputación y Fondos Europeos"
    p.font.size = Pt(9.5)
    p.font.color.rgb = WHITE

    # Pantallazo móvil del Eje a la derecha
    if os.path.exists(m_ejes):
        slide4.shapes.add_picture(m_ejes, Inches(9.5), Inches(1.7), Inches(3.0), Inches(5.3))

    # =========================================================================
    # DIAPOSITIVA 12: GAMIFICACIÓN Y LA PISCINA DE AMURJO
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, DARK_BG)
    add_slide_header(slide5, "5. Gamificación Cívica", "Premios Amurjo: La Participación Tiene Recompensa Real")

    box_amurjo = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(8.3), Inches(5.3))
    box_amurjo.fill.solid()
    box_amurjo.fill.fore_color.rgb = CARD_BG
    box_amurjo.line.color.rgb = AMURJO_CYAN

    tf_a = box_amurjo.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "LOS 4 NIVELES CÍVICOS DE ORCERA"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = AMURJO_CYAN
    p.space_after = Pt(12)

    niveles = [
        ("Nivel 1: Explorador/a del Municipio de Orcera (0-150 pts)", "Acceso libre al plan, primer voto y alta en padrón."),
        ("Nivel 2: Activista de Orcera (151-400 pts)", "Evaluación de ejes y canjes de pases para la Piscina de Amurjo."),
        ("Nivel 3: Motor de Orcera (401-800 pts)", "Prioridad en pistas de pádel, cursos DJ y talleres de robótica."),
        ("Nivel 4: Leyenda Joven de Orcera (+800 pts)", "Voz directa y debate de sus propuestas en el Pleno Municipal.")
    ]
    for n_title, n_desc in niveles:
        p = tf_a.add_paragraph()
        p.text = f"★ {n_title}"
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        p_desc = tf_a.add_paragraph()
        p_desc.text = n_desc
        p_desc.font.size = Pt(9)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(8)

    # Pantallazo móvil de Premios Amurjo a la derecha
    if os.path.exists(m_gamif):
        slide5.shapes.add_picture(m_gamif, Inches(9.5), Inches(1.7), Inches(3.0), Inches(5.3))

    # =========================================================================
    # DIAPOSITIVA 13: ASOCIACIÓN JUVENIL Y BUZÓN
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, DARK_BG)
    add_slide_header(slide6, "6. Tejido Asociativo y Escucha", "Asociación Juvenil de Orcera y Buzón Participativo")

    box_asoc = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(8.3), Inches(5.3))
    box_asoc.fill.solid()
    box_asoc.fill.fore_color.rgb = CARD_BG
    box_asoc.line.color.rgb = RGBColor(168, 85, 247)

    tf_as = box_asoc.text_frame
    tf_as.word_wrap = True
    p = tf_as.paragraphs[0]
    p.text = "REGENERANDO EL LIDERAZGO JOVEN EN EL PUEBLO"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(192, 132, 252)
    p.space_after = Pt(12)

    asoc_puntos = [
        ("Alta 100% Gratuita de Bienvenida:", "Al registrarse en el Plan, el joven recibe la membresía a cuota cero sin burocracia."),
        ("Carnet Digital Instantáneo:", "Emisión en pantalla del carnet oficial con número único (#AJO-2027-...) y código QR."),
        ("Espacio Joven en Funcionamiento:", "Organización de talleres de diseño, música, festivales culturales y excursiones."),
        ("Buzón Participativo Directo:", "Las propuestas que sumen 25 apoyos se trasladan a debate en Pleno.")
    ]
    for tit, dsc in asoc_puntos:
        p = tf_as.add_paragraph()
        p.text = f"• {tit} "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        r = p.add_run()
        r.text = dsc
        r.font.bold = False
        r.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    # Pantallazo móvil del Buzón a la derecha
    if os.path.exists(m_buzon):
        slide6.shapes.add_picture(m_buzon, Inches(9.5), Inches(1.7), Inches(3.0), Inches(5.3))

    # =========================================================================
    # DIAPOSITIVA 14: DIFUSIÓN Y ADOPCIÓN CON PANTALLAZO MÓVIL
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, DARK_BG)
    add_slide_header(slide7, "7. Cero Fricción y Difusión", "¿Cómo Llega la App a los Teléfonos de los Jóvenes?")

    canales = [
        ("1. Sin Descargas de Tiendas (PWA)", "Funciona como Aplicación Web Progresiva. No necesitan buscar en Google Play ni App Store; al abrir el enlace con 1 toque la añaden a la pantalla de inicio con icono propio.", AMURJO_CYAN),
        ("2. Difusión Viral por WhatsApp (+25 Pts)", "Cada joven dispone de su propio enlace de embajador/a. Con 1 clic en «Compartir por WhatsApp» invitan a colegas y ganan +25 Puntos Cívicos para Amurjo.", EMERALD),
        ("3. Cartelería Oficial Lista & QR Físico", "Generación e impresión directa de carteles oficiales en formato A4/A3 con código QR de alta resolución para el IES El Yelmo, bares y el Espacio Joven.", CORAL)
    ]
    for i, (c_tit, c_desc, c_col) in enumerate(canales):
        top_pos = Inches(1.7 + i * 1.78)
        box_c = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(8.3), Inches(1.6))
        box_c.fill.solid()
        box_c.fill.fore_color.rgb = CARD_BG
        box_c.line.color.rgb = c_col

        tf_c = box_c.text_frame
        tf_c.word_wrap = True
        p = tf_c.paragraphs[0]
        p.text = c_tit
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = c_col
        p.space_after = Pt(4)

        p = tf_c.add_paragraph()
        p.text = c_desc
        p.font.size = Pt(9.5)
        p.font.color.rgb = WHITE

    # Pantallazo móvil de la promoción a la derecha
    if os.path.exists(m_promo):
        slide7.shapes.add_picture(m_promo, Inches(9.5), Inches(1.7), Inches(3.0), Inches(5.3))

    # =========================================================================
    # DIAPOSITIVA 15: CONSECUENCIAS Y BENEFICIOS
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, DARK_BG)
    add_slide_header(slide8, "8. Retorno Político y Social", "Consecuencias y Beneficios para el Ayuntamiento de Orcera")

    beneficios = [
        ("🏆 Liderazgo Comarcal y Provincial", "Orcera se posiciona como el municipio pionero e innovador de la Sierra de Segura y de la provincia de Jaén en modernización de políticas de juventud.", EMERALD),
        ("💶 Captación de Subvenciones", "Disponer de métricas, datos contables en vivo y participación demostrable nos sitúa a la cabeza para captar fondos del IAJ, Diputación y Fondos Europeos.", AMURJO_CYAN),
        ("🌲 Reto Demográfico y Arraigo", "Frenamos la fuga de talento haciendo que los jóvenes se sientan protagonistas activos, escuchados y con alternativas reales de ocio y empleo en su pueblo.", CORAL),
        ("🔒 Cero Costes de Licencias", "Desarrollado con estándares web abiertos, sin pagar cuotas mensuales a multinacionales externas ni depender de empresas intermediarias.", RGBColor(168, 85, 247))
    ]

    for i, (b_tit, b_desc, b_col) in enumerate(beneficios):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.9 + row * 2.65)
        box_b = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.4))
        box_b.fill.solid()
        box_b.fill.fore_color.rgb = CARD_BG
        box_b.line.color.rgb = b_col

        tf_b = box_b.text_frame
        tf_b.word_wrap = True
        p = tf_b.paragraphs[0]
        p.text = b_tit
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = b_col
        p.space_after = Pt(8)

        p = tf_b.add_paragraph()
        p.text = b_desc
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE

    # =========================================================================
    # DIAPOSITIVA 16: PROPUESTA DE ACUERDO PLENARIO
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, DARK_BG)
    add_slide_header(slide9, "9. Votación y Aprobación", "Propuesta de Acuerdo Plenario para su Ratificación Oficial")

    box_acuerdo = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
    box_acuerdo.fill.solid()
    box_acuerdo.fill.fore_color.rgb = PINE_GREEN
    box_acuerdo.line.color.rgb = EMERALD

    tf_ac = box_acuerdo.text_frame
    tf_ac.word_wrap = True
    p = tf_ac.paragraphs[0]
    p.text = "DICTAMEN / ACUERDO DEL PLENO DEL AYUNTAMIENTO DE ORCERA:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(110, 231, 183)
    p.space_after = Pt(16)

    acuerdos = [
        "PRIMERO: Aprobar con carácter oficial el «III Plan Municipal de Juventud de Orcera (2027–2031)» como instrumento rector de las políticas locales de juventud del municipio.",
        "SEGUNDO: Ratificar la puesta en marcha de la Plataforma Web y Aplicación Cívica Municipal como canal oficial de participación, difusión, transparencia presupuestaria y dinamización juvenil.",
        "TERCERO: Respaldar la reactivación formal de la Asociación Juvenil de Orcera, incorporando la cuota cero de bienvenida y el carnet digital de socio a través de la plataforma.",
        "CUARTO: Facultar a la Alcaldía-Presidencia y a la Concejalía delegada de Juventud para la ejecución, dinamización y seguimiento de las acciones recogidas en dicho plan."
    ]
    for ac in acuerdos:
        p = tf_ac.add_paragraph()
        p.text = f"• {ac}"
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.space_after = Pt(10)

    # Guardar presentación
    out_dir = os.path.join(os.path.dirname(__file__), '..', 'docs')
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, 'Presentacion_III_Plan_Juventud_Orcera.pptx')
    prs.save(out_path)
    print(f"ÉXITO: Presentación de 16 diapositivas generada correctamente en: {out_path}")

if __name__ == '__main__':
    create_presentation()
